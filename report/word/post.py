"""Style the pandoc output to the Round-2 format: Times New Roman 12pt, single spacing, A4."""
import copy
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

src, dst = sys.argv[1], sys.argv[2]
doc = Document(src)
FONT = "Times New Roman"
BLUE = RGBColor(0x07, 0x66, 0xD1)


def set_font(rpr_owner, name=FONT):
    rpr = rpr_owner.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in list(rf.attrib):
        del rf.attrib[a]
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), name)


# ---- page setup ------------------------------------------------------------------------
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(2.2))
sec.header_distance = Cm(1.2); sec.footer_distance = Cm(1.2); sec.gutter = Cm(0)
TEXT_W = sec.page_width - sec.left_margin - sec.right_margin

S = lambda n: next(x for x in doc.styles if x.name == n)
# ---- styles ------------------------------------------------------------------------------
# docDefaults font
rpr_default = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
rf = rpr_default.find(qn("w:rFonts"))
if rf is None:
    rf = OxmlElement("w:rFonts"); rpr_default.insert(0, rf)
for a in list(rf.attrib):
    del rf.attrib[a]
for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
    rf.set(qn(a), FONT)

for st in doc.styles:
    if st.type != 1 and st.type != 2:  # paragraph / character
        continue
    set_font(st.element)
    if st.type == 1:
        pf = st.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE

for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text", "Bibliography"):
    if name in [s.name for s in doc.styles]:
        st = S(name)
        st.font.size = Pt(12)
        st.paragraph_format.space_before = Pt(0)
        st.paragraph_format.space_after = Pt(4)
        if name in ("Body Text", "First Paragraph"):
            st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
S("Compact").paragraph_format.space_after = Pt(1)
bt = S("Block Text")
bt.paragraph_format.left_indent = Cm(0.6); bt.paragraph_format.right_indent = Cm(0.6)

heads = {"Heading 1": (15, True, False, BLUE, 14, 6), "Heading 2": (12.5, True, False, None, 10, 4),
         "Heading 3": (12, True, True, None, 8, 3)}
for name, (size, bold, ital, color, before, after) in heads.items():
    st = S(name)
    st.font.size = Pt(size); st.font.bold = bold; st.font.italic = ital
    st.font.color.rgb = color if color else RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for name in ("Image Caption", "Table Caption", "Caption"):
    if name in [s.name for s in doc.styles]:
        st = S(name)
        st.font.size = Pt(10); st.font.italic = False
        st.paragraph_format.space_before = Pt(2); st.paragraph_format.space_after = Pt(8)
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
S("Table Caption").paragraph_format.keep_with_next = True

for name in ("Source Code", "Verbatim Char"):
    if name in [s.name for s in doc.styles]:
        set_font(S(name).element, "Courier New")
        S(name).font.size = Pt(8)

# ---- bold "Figure N:" / "Table N:" prefixes in captions ----------------------------------
cap_re = re.compile(r"^((Figure|Table) \d+:)")
for p in doc.paragraphs:
    if p.style.name in ("Image Caption", "Table Caption", "Caption") and p.runs:
        r0 = p.runs[0]
        m = cap_re.match(r0.text)
        if m:
            rest = r0.text[m.end():]
            r0.text = m.group(1); r0.bold = True
            new = copy.deepcopy(r0._r); r0._r.addnext(new)
            from docx.text.run import Run
            nr = Run(new, p); nr.text = rest; nr.bold = None

# ---- unwrap pandoc "FigureTable" layout tables (figure pairs) into one paragraph -----------
for tbl in list(doc.element.body.iter(qn("w:tbl"))):
    st = tbl.find(qn("w:tblPr")).find(qn("w:tblStyle"))
    if st is None or st.get(qn("w:val")) != "FigureTable":
        continue
    runs = [r for r in tbl.iter(qn("w:r")) if r.find(".//" + qn("wp:inline")) is not None]
    p = doc.add_paragraph()._p
    for r in runs:
        p.append(r)
    tbl.addprevious(p); tbl.getparent().remove(tbl)

# ---- tables: booktabs-like rules, 10pt text ----------------------------------------------
def border(tag, sz):
    el = OxmlElement(f"w:{tag}")
    el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(sz)); el.set(qn("w:space"), "0"); el.set(qn("w:color"), "000000")
    return el

for t in doc.tables:
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = t._tbl.tblPr
    old = tblPr.find(qn("w:tblBorders"))
    if old is not None:
        tblPr.remove(old)
    b = OxmlElement("w:tblBorders"); b.append(border("top", 12)); b.append(border("bottom", 12))
    after = [tblPr.find(qn(t)) for t in ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription")]
    after = [a for a in after if a is not None]
    if after:
        after[0].addprevious(b)
    else:
        tblPr.append(b)
    tw = tblPr.find(qn("w:tblW"))
    if tw is None:
        tw = OxmlElement("w:tblW"); tblPr.append(tw)
    tw.set(qn("w:type"), "pct"); tw.set(qn("w:w"), "5000")
    for ri, row in enumerate(t.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1); p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.size = Pt(10)
            if ri == 0:
                tcPr = cell._tc.get_or_add_tcPr()
                tb = OxmlElement("w:tcBorders"); tb.append(border("bottom", 6)); tcPr.append(tb)
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.bold = True

# space after each table
for t in doc.tables:
    nxt = t._tbl.getnext()
    if nxt is not None and nxt.tag == qn("w:p"):
        from docx.text.paragraph import Paragraph
        Paragraph(nxt, None).paragraph_format.space_before = Pt(8)

# merge consecutive image-only paragraphs (figure pairs) into one paragraph
def img_only(el):
    return el is not None and el.tag == qn("w:p") and el.findall(".//" + qn("wp:inline")) and not "".join(t.text or "" for t in el.iter(qn("w:t"))).strip()
for p in list(doc.paragraphs):
    el = p._p
    if el.getparent() is None or not img_only(el):
        continue
    nxt = el.getnext()
    while img_only(nxt):
        for r in nxt.findall(qn("w:r")):
            el.append(r)
        after = nxt.getnext(); nxt.getparent().remove(nxt); nxt = after

# ---- images: fit to text width; two-up figures side by side -------------------------------
for p in doc.paragraphs:
    inl = p._p.findall(".//" + qn("wp:inline"))
    if not inl:
        continue
    maxw = int(TEXT_W * (0.45 if len(inl) > 1 else 0.92))
    for i in inl:
        ext = i.find(qn("wp:extent"))
        cx, cy = int(ext.get("cx")), int(ext.get("cy"))
        scale = maxw / cx
        ncx, ncy = int(cx * scale), int(cy * scale)
        maxh = int(Cm(15))
        if ncy > maxh:
            ncx, ncy = int(ncx * maxh / ncy), maxh
        ext.set("cx", str(ncx)); ext.set("cy", str(ncy))
        for e in i.iter(qn("a:ext")):
            e.set("cx", str(ncx)); e.set("cy", str(ncy))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True

# ---- title page + TOC ---------------------------------------------------------------------
body = doc.element.body
first = body[0]


def para_before(anchor, text="", size=12, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=0):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    if text:
        r = p.add_run(text); r.font.size = Pt(size); r.bold = bold
        if color:
            r.font.color.rgb = color
    anchor.addprevious(p._p)
    return p


para_before(first, "D2S Bharat", 30, True, BLUE, before=110, after=10)
para_before(first, "From Data-Science Job Demand to Skills, Success and Training Decisions", 16, after=30)
para_before(first, "Analytics approach note — Round 2", 14)
para_before(first, "SAS CU Hackathon · Build for Bharat", 14, after=50)
para_before(first, "Team Code5urge", 14, True, after=6)
para_before(first, "Shaman Sharma    Anurag Prajapati    Amrita Raman    Rhydham Mittal", 13, after=160)
para_before(first, "Data: SAS-provided Analytics Jobs, DataScience Jobs, JDS Skill Traits and SDS Personality "
                   "Traits files. All code and outputs are reproducible from the project repository (Appendix A).", 10)
pb = para_before(first)
pb.add_run().add_break(WD_BREAK.PAGE)

toc_h = para_before(first, "Contents", 15, True, BLUE, align=WD_ALIGN_PARAGRAPH.LEFT, after=6)
tp = para_before(first, align=WD_ALIGN_PARAGRAPH.LEFT)
r = tp.add_run()
for kind, text in (("begin", None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None)):
    if kind:
        fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); r._r.append(fc)
    else:
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = text; r._r.append(it)
r2 = tp.add_run("Right-click and choose “Update Field” to build the table of contents.")
r2.italic = True
r3 = tp.add_run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "end"); r3._r.append(fc)
pb = para_before(first)
pb.add_run().add_break(WD_BREAK.PAGE)

# Page break before each top-level numbered section that the PDF starts on a new page
for p in doc.paragraphs:
    if p.style.name == "Heading 1" and (p.text.startswith("1 ") or p.text.startswith("References")
                                        or p.text.startswith("Appendix A")):
        p.paragraph_format.page_break_before = True

# update fields on open so the TOC fills in
settings = doc.settings.element
uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true"); settings.append(uf)

# ---- footer page numbers ------------------------------------------------------------------
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run()
for kind, text in (("begin", None), (None, "PAGE"), ("separate", None), (None, None), ("end", None)):
    if kind:
        fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); r._r.append(fc)
    elif text:
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = text; r._r.append(it)
    else:
        t = OxmlElement("w:t"); t.text = "1"; r._r.append(t)
r.font.size = Pt(10)
sec.different_first_page_header_footer = True

doc.core_properties.title = "D2S Bharat — Analytics approach note (Round 2)"
doc.core_properties.author = "Team Code5urge"
doc.save(dst)
print("saved", dst)
